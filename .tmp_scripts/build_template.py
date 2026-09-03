import sys
from docx import Document
from docx.oxml.ns import qn

LABEL_MAP = {
    "Name (中文名称)": "party_a_name_cn",
    "Name (英文名称)": "party_a_name_en",
    "Add(中文地址)": "party_a_address_cn",
    "Add(英文地址)": "party_a_address_en",
    "Zip Code(邮编)": "party_a_zip",
    "Contact Person": "party_a_contact",
    "Tel(": "party_a_tel",
    "E-mail (邮箱)": "party_a_email",
}


def clear_run_format(run):
    rpr = run._element.find(qn('w:rPr'))
    if rpr is None:
        return
    for tag in ('w:highlight', 'w:shd'):
        el = rpr.find(qn(tag))
        if el is not None:
            rpr.remove(el)


def set_placeholder(paragraph, placeholder):
    runs = paragraph.runs
    target_runs = [r for r in runs if is_run_highlighted(r)]
    if not target_runs:
        return False
    target_runs[0].text = placeholder
    clear_run_format(target_runs[0])
    for r in target_runs[1:]:
        r.text = ''
        clear_run_format(r)
    return True


def is_run_highlighted(r):
    rpr = r._element.find(qn('w:rPr'))
    if rpr is None:
        return False
    if rpr.find(qn('w:highlight')) is not None:
        return True
    shd = rpr.find(qn('w:shd'))
    if shd is not None and (shd.get(qn('w:fill')) or '').lower() not in ('auto', 'ffffff', ''):
        return True
    return False


def set_placeholder_range(paragraph, placeholder):
    """Replace the whole span from the first to the last highlighted run
    (inclusive of any non-highlighted separator runs in between) with a
    single placeholder. Useful for fragmented date ranges."""
    runs = paragraph.runs
    hl_idxs = [i for i, r in enumerate(runs) if is_run_highlighted(r)]
    if not hl_idxs:
        return False
    start, end = hl_idxs[0], hl_idxs[-1]
    runs[start].text = placeholder
    clear_run_format(runs[start])
    for r in runs[start + 1:end + 1]:
        r.text = ''
        clear_run_format(r)
    return True


def strip_leading_char_after_range(paragraph, char):
    """After set_placeholder_range collapses a highlighted date span into a
    single placeholder, the run immediately following it may still start
    with a trailing date-suffix character (e.g. '日') that belongs to the
    end of the collapsed range but wasn't itself highlighted. Strip it."""
    runs = paragraph.runs
    for i, r in enumerate(runs):
        if '{{' in r.text and '}}' in r.text:
            j = i + 1
            while j < len(runs) and runs[j].text == '':
                j += 1
            if j < len(runs) and runs[j].text.startswith(char):
                runs[j].text = runs[j].text[len(char):]
            break


def is_label_match(text, label):
    return text.strip().startswith(label)


def process_label_table(doc):
    for t in doc.tables:
        for row in t.rows:
            if len(row.cells) < 2:
                continue
            label_text = row.cells[0].text
            for label, key in LABEL_MAP.items():
                if is_label_match(label_text, label):
                    value_cell = row.cells[1]
                    for p in value_cell.paragraphs:
                        set_placeholder(p, "{{ " + key + " }}")
                    break


def process_special_paragraphs(doc):
    for p in doc.paragraphs:
        text = p.text
        if text.startswith("Agreement Number"):
            set_placeholder_range(p, "{{ agreement_number }}")
        elif text.startswith("This agreement will be valid from"):
            set_placeholder_range(p, " {{ effective_range_en }}")
        elif "此协议有效期自" in text or ("Part A could choose to renew" in text):
            set_placeholder_range(p, "{{ effective_range_cn }}")
            strip_leading_char_after_range(p, "日")
        elif text.startswith("PARTY A:"):
            set_placeholder(p, "{{ party_a_name_en }}")
        elif "Date" in text[:6] and ("日期" in text):
            set_placeholder(p, "{{ signing_date }}")


def process_shop_table(doc):
    for t in doc.tables:
        if len(t.rows) < 3:
            continue
        header_cells = [c.text.strip() for c in t.rows[1].cells] if len(t.rows) > 1 else []
        if not header_cells or '销售平台' not in header_cells[0]:
            continue
        data_row_idx = 2
        row = t.rows[data_row_idx]
        placeholders = [
            "{%tr for s in shops %}{{ s.platform }}",
            "{{ s.shop_url }}",
            "{{ s.shop_id }}",
            "{{ s.shop_name }}",
            "{{ s.brand_names }}",
            "{{ s.product_names_cn }}",
            "{{ s.product_names_en }}{%tr endfor %}",
        ]
        for cell, ph in zip(row.cells, placeholders):
            for p in cell.paragraphs:
                for r in list(p.runs):
                    r.text = ''
                    clear_run_format(r)
                if p.runs:
                    p.runs[0].text = ph
                else:
                    p.add_run(ph)
        extra_rows = []
        for r in t.rows[data_row_idx + 1:]:
            texts = [c.text.strip() for c in r.cells]
            if any(texts):
                extra_rows.append(r)
        for r in extra_rows:
            r._element.getparent().remove(r._element)
        return True
    return False


def main():
    src, dst = sys.argv[1], sys.argv[2]
    doc = Document(src)
    process_label_table(doc)
    process_special_paragraphs(doc)
    ok = process_shop_table(doc)
    doc.save(dst)
    print("shop table processed:", ok)


if __name__ == "__main__":
    main()
