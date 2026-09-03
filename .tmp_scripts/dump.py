import sys
from docx import Document

path = sys.argv[1]
doc = Document(path)

print("--- paragraphs with braces ---")
for idx, p in enumerate(doc.paragraphs):
    if '{{' in p.text or '{%' in p.text:
        print(idx, repr(p.text))

print("--- tables ---")
for ti, t in enumerate(doc.tables):
    print(f"table {ti} rows={len(t.rows)}")
    for ri, row in enumerate(t.rows):
        cells = [c.text for c in row.cells]
        if any('{{' in c or '{%' in c for c in cells) or ti == 0:
            print(ri, cells)
